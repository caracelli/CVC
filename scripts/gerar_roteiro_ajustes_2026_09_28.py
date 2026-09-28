# -*- coding: utf-8 -*-
"""Gera ENTREGA/ROTEIRO_AJUSTES_CVC_IAM_2026-09-28.docx (+ .md).

Responde os SEIS pontos do documento "ajustes_apl_28_09.docx" (retorno da
Bruna ao pacote de 24/09), cada um validado visualmente pelo usuario no painel
em 28/09, e avisa o que muda na tela.

Mesma formatacao dos roteiros de 08, 11, 18 e 24/09.

OS NUMEROS SAO DA BASE DELA: a de 15/09, com os ultimos SYSTUR, Oracle e RH
devolvidos a ENTRADA (para serem relidos com a leitura de acentos corrigida,
como a proxima execucao dela fara') e reprocessada com o codigo desta entrega;
comparada contra o banco do pacote de 24/09.

Uso:  python scripts/gerar_roteiro_ajustes_2026_09_28.py
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import gerar_roteiro_ajustes as base  # noqa: E402

from docx import Document  # noqa: E402
from docx.shared import Pt  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402

h1, par, nota = base.h1, base.par, base.nota
AZUL, CINZA, TEXTO = base.AZUL, base.CINZA, base.TEXTO
_MD = base._MD
MEDIDO = "Medido na sua base (a de 15/09, reprocessada aqui com esta versão)"

RAIZ = Path(__file__).resolve().parent.parent
OUT_DOCX = RAIZ / "ENTREGA" / "ROTEIRO_AJUSTES_CVC_IAM_2026-09-28.docx"
OUT_MD = RAIZ / "ENTREGA" / "ROTEIRO_AJUSTES_CVC_IAM_2026-09-28.md"


def regra(*a, **kw):
    kw.setdefault("rot_medido", MEDIDO)
    return base.regra(*a, **kw)


def lista(doc, itens):
    for txt in itens:
        p = doc.add_paragraph(txt, style="List Bullet")
        for r_ in p.runs:
            r_.font.size = Pt(9.5)
            r_.font.color.rgb = TEXTO
        _MD.append(f"- {txt}")
    _MD.append("")


def main():
    _MD.clear()
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)

    # ------------------------------------------------------------- capa
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = t.add_run("CVC IAM Analytics")
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = AZUL
    s = doc.add_paragraph()
    r = s.add_run("Roteiro de validação — seu documento de 28/09")
    r.font.size = Pt(13)
    r.font.color.rgb = CINZA
    par(doc, "28/09/2026", size=9, cor=CINZA, md=False)
    _MD.insert(0, "# CVC IAM Analytics\n"
                  "## Roteiro de validação — seu documento de 28/09\n"
                  "28/09/2026\n")

    doc.add_paragraph()
    par(doc, "Por que este documento existe", bold=True, size=11, cor=AZUL, space=2)
    par(doc,
        "Ele responde, um a um, os seis pontos do seu documento de 28/09 "
        "(ajustes_apl_28_09). Todos precisavam de correção. Cada item diz o que "
        "era, o que mudou, como conferir e o valor esperado na sua base.")
    par(doc, "Leia primeiro: as pendências DIMINUEM", bold=True, size=11,
        cor=AZUL, space=2)
    par(doc,
        "As pendências caem de 1.278 para 1.072 pessoas, porque quem tem no SIG "
        "exatamente os perfis da função deixou de ser pendência (item 5). O total "
        "de linhas sobe de 12.310 para 13.446, porque o Oracle previsto voltou a "
        "aparecer para 33 pessoas que não têm SYSTUR (item 4).")
    nota(doc,
         "O pacote já traz a pasta DADOS com o banco processado: você NÃO "
         "precisa rodar o Processador. Se rodar, tudo bem — deixe as matrizes na "
         "pasta ENTRADA, como da última vez.")

    # ------------------------------------------------- antes de tudo
    doc.add_page_break()
    h1(doc, "Antes de tudo: o que fazer, nesta ordem")
    lista(doc, [
        "1. Feche o painel e o Processador, se estiverem abertos.",
        "2. BACKUP (obrigatório): copie as pastas DADOS\\BANCO e INTERACOES "
        "para outro lugar. Se algo não sair como esperado, é só devolvê-las.",
        "3. Extraia o pacote na pasta principal da instalação — a que contém "
        "DADOS, ENTRADA e EXECUTAVEIS — e aceite substituir os arquivos.",
        "4. NÃO apague nem mexa na pasta INTERACOES: é onde ficam as suas "
        "tratativas, e o painel as lê ao vivo.",
        "5. Abra o visualizador.exe.",
    ])

    # ------------------------------------------------- os 6 pontos
    doc.add_page_break()
    h1(doc, "1. Os seis pontos do seu documento de 28/09")

    regra(doc, "1", "O Excel volta a bater com a tela",
          prioritaria=True,
          decide="O que sai na planilha da Consulta no formato Analítico.",
          criterio="No Analítico, cada linha de sistema completava o que vinha "
                   "vazio com o valor da linha da pessoa — e a linha da pessoa "
                   "junta os perfis de todos os sistemas. Por isso os perfis do "
                   "SIG apareciam como “Perfil Encontrado” em todos os sistemas. "
                   "Agora só a identificação (matrícula, nome, cargo, centro de "
                   "custo, e-mail) se repete; perfil vazio num sistema quer dizer "
                   "que a pessoa não tem acesso nele. O mesmo defeito fazia as "
                   "exportações de Inclusão, Histórico RH e Quarentena perderem o "
                   "primeiro registro de cada pessoa no Analítico — corrigido junto.",
          conferir="Consulta → busque BRUNA OLIVEIRA SANTOS → Exportar → "
                   "Analítico.",
          esperado="Sete linhas com a matrícula PREST-corpp03779. “Perfil "
                   "Encontrado” preenchido só na linha do SIG; nas demais, vazio. "
                   "“Perfil Esperado” com o perfil de cada sistema.")

    regra(doc, "2", "Oracle: o mesmo perfil não aparece mais em “falta” e “a mais”",
          prioritaria=True,
          decide="Como a aplicação lê os acentos dos arquivos.",
          criterio="Os seus arquivos estão certos. A aplicação tentava adivinhar "
                   "a codificação e às vezes errava: “Criação” virava “Criaçăo” e "
                   "não batia com a matriz. A leitura foi corrigida para TODOS os "
                   "arquivos. Na sua base isso atingia 47 perfis do Oracle, 110 "
                   "nomes do SYSTUR e departamento/cargo do RH — e a data de "
                   "admissão, que vinha vazia. Correções só de acento não são "
                   "tratadas como movimentação: o Histórico e os Transferidos "
                   "não mudam.",
          conferir="Consulta → HERMES LUIS DOS SANTOS (90000639) → Oracle EBS.",
          esperado="“Tem 46 · Faltam 3: CVC AR RA VIAGENS Consulta Fiscal, CVC PA "
                   "NOVA VISUAL Analista, CVC RI NOVA VISUAL Consulta Fiscal · 1 a "
                   "mais: CVC OIE BRASIL - Relatório de Despesas” — exatamente a "
                   "sua conferência manual. Os “CVC INV … Criação Itens” estão "
                   "entre os 46 que ele tem. Na base: 42 falsos “perfil inválido” "
                   "do Oracle deixaram de existir.")

    regra(doc, "3", "A lista de perfis sai sempre do mesmo jeito",
          decide="Como a tela mostra vários perfis de um mesmo sistema.",
          criterio="Quando cada perfil era uma linha separada, a tela os juntava "
                   "com vírgula; quando vinham juntos, mostrava em lista. Agora é "
                   "sempre um embaixo do outro, com os 6 primeiros visíveis e "
                   "“+N outros” para abrir o resto — com a mesma formatação.",
          conferir="Consulta → LETICIA LEMOS DE SOUZA (34532586) → Oracle EBS → "
                   "clique em “+17 outros”.",
          esperado="Os 23 acessos esperados em lista vertical, do começo ao fim.")

    regra(doc, "4", "O Oracle previsto aparece para quem não tem SYSTUR",
          prioritaria=True,
          decide="O que a aplicação mostra no Oracle para quem não tem perfil "
                 "no SYSTUR nem acesso no Oracle.",
          criterio="A regra do Oracle segue o perfil do SYSTUR. Sem perfil no "
                   "SYSTUR, ela não tinha com o que comparar e escondia todo o "
                   "Oracle — a tela dizia “sem perfil previsto”, o que não é "
                   "verdade. Agora vale o perfil de SYSTUR que a matriz prevê para "
                   "a pessoa (o mesmo que a pendência de SYSTUR manda incluir). "
                   "Quem TEM Oracle e não tem SYSTUR continua como estava: a "
                   "pendência fica no SYSTUR.",
          conferir="Consulta → RAFAEL FELIPE DE MORAES (14546).",
          esperado="Oracle EBS em “Acessos esperados” com 39 perfis, junto de "
                   "SIGOT (Contabil1) e SYSTUR (INTEGRADOR_CONTABIL). O Oracle sai "
                   "do bloco “Sem mapeamento”. Na base: 33 pessoas.")

    regra(doc, "5", "SIG com exatamente os perfis da função é aderente",
          prioritaria=True,
          decide="Se quem tem no SIG o conjunto previsto pela função vira "
                 "pendência por “mais de um perfil”.",
          criterio="Sua resposta à pergunta de 24/09: no SIG os perfis da função "
                   "se somam. Quem tem exatamente o conjunto previsto passa a ser "
                   "aderente, sem o alerta, e a lista sai uma vez só (não mais "
                   "“Tem hoje” + “Deveria ter” iguais). Quem tem só parte do "
                   "conjunto, ou algo a mais, continua em análise. Nos outros "
                   "sistemas a regra de “um perfil por sistema” não muda.",
          conferir="Consulta → ATHAMIRIS DA SILVA TORRES (23242).",
          esperado="SIG em “Acessos encontrados”, em verde, sem alerta. Em "
                   "“Funções previstas”, Supervisor de Operações “completa”. Na "
                   "base: 219 linhas do SIG deixaram de ser pendência.")

    regra(doc, "6", "As funções mostram um perfil por linha",
          decide="Como aparece o sistema com vários perfis dentro de “Funções "
                 "previstas”.",
          criterio="O SIG vinha num parágrafo com um status só (“Em Análise”), "
                   "enquanto “Outras funções” mostrava perfil a perfil. Agora as "
                   "duas têm a mesma leitura: cada perfil previsto aparece como "
                   "“tem” ou “falta”, o que a pessoa tem fora do previsto como “a "
                   "mais”, e a conta da função é por perfil.",
          conferir="Consulta → PAMELLA DE OLIVEIRA DA SILVA (6171) → Funções "
                   "previstas → clique em “Operacional e Não Operacional”.",
          esperado="“completa”, com cada perfil do SIG numa linha marcada “tem”, "
                   "seguidos de SIGOT e SYSTUR (“tem”) e Opera (“sem extrato”).")

    # ------------------------------------------------- o que muda na tela
    doc.add_page_break()
    h1(doc, "2. O que vai parecer diferente — e por quê")
    par(doc,
        "Estas mudanças alteram números que você acompanha. Nenhuma delas é "
        "erro; todas vêm dos pontos acima.")
    lista(doc, [
        "As pendências CAEM de 1.278 para 1.072 pessoas (de 1.342 para 1.123 "
        "linhas): são as 219 linhas do SIG com o conjunto exato da função (item 5).",
        "O total de linhas SOBE de 12.310 para 13.446: são os perfis de Oracle a "
        "incluir das 33 pessoas sem SYSTUR (item 4).",
        "A data de admissão passa a aparecer. Nesta base, 1.997 pessoas (as do "
        "último arquivo de RH); as demais entram à medida que os próximos "
        "arquivos chegarem.",
        "Histórico e Transferidos não mudam: correções só de acento não contam "
        "como movimentação (item 2).",
    ])

    # ------------------------------------------------- registro
    doc.add_page_break()
    h1(doc, "3. Suas respostas às perguntas de 24/09")
    par(doc,
        "4.1 — Relatório de Despesas: você o apontou como “a mais” no Oracle do "
        "HERMES. Continua contando como pendência, como já estava.")
    par(doc,
        "4.2 — SIG: os perfis da função se somam. Aplicado no item 5.")

    # ------------------------------------------------- saida
    OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_DOCX)
    OUT_MD.write_text("\n".join(_MD), encoding="utf-8")
    print(f"gerado: {OUT_DOCX}")
    print(f"gerado: {OUT_MD}")


if __name__ == "__main__":
    main()
