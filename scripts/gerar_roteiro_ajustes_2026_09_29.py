# -*- coding: utf-8 -*-
"""Gera ENTREGA/ROTEIRO_AJUSTES_CVC_IAM_2026-09-29.docx (+ .md).

Responde o documento "ajustes_apl_29_09.pdf" (retorno da Bruna): na CCO, a
funcao da pessoa passa a ser identificada pelo perfil que ela tem em SYSTUR,
SICA RA, SICA ESFERA e SIGOT (validacao reversa pela CCO); Oracle e SIG so'
recebem o que a funcao preve. Validado visualmente pelo usuario em 29/09.

Traz os 19 casos que mudaram na base dela, com antes/depois, para ela conferir.

Mesma formatacao dos roteiros de 08, 11, 18, 24 e 28/09.

OS NUMEROS SAO DA BASE DELA: a mesma do pacote de 28/09, reprocessada com o
codigo desta entrega e comparada contra o banco daquele pacote.

Uso:  python scripts/gerar_roteiro_ajustes_2026_09_29.py
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import gerar_roteiro_ajustes as base  # noqa: E402

from docx import Document  # noqa: E402
from docx.shared import Pt, Cm  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402

h1, par, nota = base.h1, base.par, base.nota
AZUL, CINZA, TEXTO = base.AZUL, base.CINZA, base.TEXTO
_MD = base._MD
MEDIDO = "Medido na sua base (a do pacote de 28/09, reprocessada com esta versão)"

RAIZ = Path(__file__).resolve().parent.parent
OUT_DOCX = RAIZ / "ENTREGA" / "ROTEIRO_AJUSTES_CVC_IAM_2026-09-29.docx"
OUT_MD = RAIZ / "ENTREGA" / "ROTEIRO_AJUSTES_CVC_IAM_2026-09-29.md"

# Os 19 casos que mudaram (base de 15/09 reprocessada, 29/09/2026):
# (matricula, nome, perfil que identifica, funcao, antes -> depois)
CASOS = [
    ("34532401", "BRUNA OLIVEIRA FERREIRA DA SILVA", "SICA RA POS FATURAMENTO CONC",
     "Pós Faturamento", "5 funções cobradas, 12 esperados", "1 função, 0 esperados"),
    ("34532403", "MICHELE MARIANO DA SILVA", "SICA RA POS FAT TERRESTE CASH",
     "Pós Faturamento Terreste_Cash", "6 funções, 12 esperados", "1 função, 0 esperados"),
    ("34532737", "ANA CAROLINE JARDIM BELLO", "SICA RA POS FAT ESFERA",
     "Pós Faturamento - Esfera", "5 funções, 12 esperados", "1 função, 1 esperado"),
    ("34532404", "PALOMA CAMPOREZE DE SOUZA", "SICA RA POS FATURAMENTO I",
     "Pós Faturamento I", "5 funções, 12 esperados", "1 função, 1 esperado"),
    ("34532734", "ALYNE VIANA SOUZA ROCHA", "SIGOT POSTERRESTRE",
     "Pós Terrestre", "6 funções, 10 esperados", "1 função, 0 esperados"),
    ("34532396", "JESSICA LOPES OLIVEIRA", "SIGOT POSTERRESTRE",
     "Pós Terrestre", "6 funções, 10 esperados", "1 função, 0 esperados"),
    ("34532735", "PALOMA ISABEL DA CUNHA QUEIROZ", "SIGOT POSTERRESTRE",
     "Pós Terrestre", "6 funções, 10 esperados", "1 função, 0 esperados"),
    ("34532321", "JESSICA LUZIA VASCO SIMOES", "SIGOT e SICA RA CP_BACKOFFICE",
     "CP BACKOFFICE", "7 funções, 12 esperados, 1 pendência", "1 função, 2 esperados, 0 pendência"),
    ("90000054", "ADRIANA MARQUES DE SOUZA SOARES", "SIGOT e SICA RA CP_BACKOFFICE",
     "CP BACKOFFICE", "7 funções, 12 esperados, 1 pendência", "1 função, 2 esperados, 0 pendência"),
    ("12698", "ERICKA CRISTINA REIS", "SIGOT Carros_Hoteis_CP",
     "Carros e Hoteis", "7 funções, 12 esperados, 1 pendência", "1 função, 3 esperados, 0 pendência"),
    ("15161", "JESSICA MACHADO DOS REIS", "SIGOT Carros_Hoteis_CP",
     "Carros e Hoteis", "7 funções, 12 esperados, 1 pendência", "1 função, 3 esperados, 0 pendência"),
    ("34532590", "NATALI ALMEIDA DE OLIVEIRA", "SIGOT Carros_Hoteis_CP",
     "Carros e Hoteis", "7 funções, 12 esperados, 1 pendência", "1 função, 3 esperados, 0 pendência"),
    ("90000066", "CILENE DO NASCIMENTO CRUZ", "SIGOT Carros_Hoteis_CP",
     "Carros e Hoteis", "7 funções, 12 esperados, 1 pendência", "1 função, 3 esperados, 0 pendência"),
    ("33057", "WAGNER NOVAES FAGUNDES", "SIGOT Carros_Hoteis_CP",
     "Carros e Hoteis", "7 funções, 12 esperados, 1 pendência", "1 função, 3 esperados, 0 pendência"),
    ("34532095", "KATHELIN LOPES CORREA", "SIGOT Internacional_CP",
     "Internacional", "8 funções, 12 esperados, 8 pendências", "1 função, 1 esperado, 1 pendência"),
    ("34532254", "JESSICA ALVES MUNHOZ", "SIGOT PGTOS_NACIONAIS_CP_I",
     "Pagamentos Nacionais I", "8 funções, 8 esperados, 1 pendência", "1 função, 1 esperado, 0 pendência"),
    ("90000177", "HALANA TAROSSI NASCIMENTO", "SIGOT Atd_For_TREND_N2",
     "Atendimento a fornecedores TREND N2", "7 funções, 6 esperados", "1 função, 1 esperado"),
    ("90001090", "WELLINGTON RODRIGUES DE OLIVEIRA", "SIGOT SupFin",
     "Sup Financeiro", "3 funções, 3 esperados", "1 função, 0 esperados"),
    ("34531885", "CAROLINE AGUILAR DE OLIVEIRA", "SIGOT SupFin",
     "Sup Financeiro", "3 funções, 3 esperados", "1 função, 0 esperados"),
]

# Antes/depois de cada caso, lidos dos dois bancos (pacote 28/09 x 29/09).
ANTES_DEPOIS = {
    '34532401': ('Funções: Pos A Receber Comissão, Pós Faturamento, Pós Faturamento - Esfera, Pós Faturamento I, Pós Terrestre. Bloco “Acessos esperados”: SICA_ESFERA 3, SIG 5, SIGOT 2, SYSTUR 2. Pendências: nenhuma.',
                 'Função: Pós Faturamento. Bloco “Acessos esperados”: nenhum. Pendências: nenhuma.'),
    '34532403': ('Funções: Pos A Receber Comissão, Pós Faturamento, Pós Faturamento - Esfera, Pós Faturamento I, Pós Faturamento Terreste_Cash, Pós Terrestre. Bloco “Acessos esperados”: SICA_ESFERA 3, SIG 5, SIGOT 2, SYSTUR 2. Pendências: nenhuma.',
                 'Função: Pós Faturamento Terreste_Cash. Bloco “Acessos esperados”: nenhum. Pendências: nenhuma.'),
    '34532737': ('Funções: Pos A Receber Comissão, Pós Faturamento, Pós Faturamento - Esfera, Pós Faturamento I, Pós Terrestre. Bloco “Acessos esperados”: SICA_ESFERA 3, SIG 5, SIGOT 2, SYSTUR 2. Pendências: nenhuma.',
                 'Função: Pós Faturamento - Esfera. Bloco “Acessos esperados”: SICA_ESFERA 1. Pendências: nenhuma.'),
    '34532404': ('Funções: Pos A Receber Comissão, Pós Faturamento, Pós Faturamento - Esfera, Pós Faturamento I, Pós Terrestre. Bloco “Acessos esperados”: SICA_ESFERA 3, SIG 5, SIGOT 2, SYSTUR 2. Pendências: nenhuma.',
                 'Função: Pós Faturamento I. Bloco “Acessos esperados”: SICA_ESFERA 1. Pendências: nenhuma.'),
    '34532734': ('Funções: Pos A Receber Comissão, Pós Faturamento, Pós Faturamento - Esfera, Pós Faturamento I, Pós Faturamento Terreste_Cash, Pós Terrestre. Bloco “Acessos esperados”: SICA_ESFERA 3, SICA_RA 5, SYSTUR 2. Pendências: nenhuma.',
                 'Função: Pós Terrestre. Bloco “Acessos esperados”: nenhum. Pendências: nenhuma.'),
    '34532396': ('Funções: Pos A Receber Comissão, Pós Faturamento, Pós Faturamento - Esfera, Pós Faturamento I, Pós Faturamento Terreste_Cash, Pós Terrestre. Bloco “Acessos esperados”: SICA_ESFERA 3, SICA_RA 5, SYSTUR 2. Pendências: nenhuma.',
                 'Função: Pós Terrestre. Bloco “Acessos esperados”: nenhum. Pendências: nenhuma.'),
    '34532735': ('Funções: Pos A Receber Comissão, Pós Faturamento, Pós Faturamento - Esfera, Pós Faturamento I, Pós Faturamento Terreste_Cash, Pós Terrestre. Bloco “Acessos esperados”: SICA_ESFERA 3, SICA_RA 5, SYSTUR 2. Pendências: nenhuma.',
                 'Função: Pós Terrestre. Bloco “Acessos esperados”: nenhum. Pendências: nenhuma.'),
    '34532321': ('Funções: CP BACKOFFICE, Carros e Hoteis, Coordenação Suporte + Carro e Hotéis, Internacional, Não Operacional, Suporte N1, Suporte N2. Bloco “Acessos esperados”: SICA_ESFERA 2, SICA_RA 4, SYSTUR 6. Pendências: SIG 1.',
                 'Função: CP BACKOFFICE. Bloco “Acessos esperados”: SICA_ESFERA 1, SICA_RA 1. Pendências: nenhuma.'),
    '90000054': ('Funções: CP BACKOFFICE, Carros e Hoteis, Coordenação Suporte + Carro e Hotéis, Internacional, Não Operacional, Suporte N1, Suporte N2. Bloco “Acessos esperados”: SICA_ESFERA 2, SICA_RA 4, SYSTUR 6. Pendências: SIG 1.',
                 'Função: CP BACKOFFICE. Bloco “Acessos esperados”: SICA_ESFERA 1, SICA_RA 1. Pendências: nenhuma.'),
    '12698': ('Funções: CP BACKOFFICE, Carros e Hoteis, Coordenação Suporte + Carro e Hotéis, Internacional, Não Operacional, Suporte N1, Suporte N2. Bloco “Acessos esperados”: SICA_ESFERA 2, SICA_RA 4, SYSTUR 6. Pendências: SIG 1.',
                 'Função: Carros e Hoteis. Bloco “Acessos esperados”: SICA_ESFERA 1, SICA_RA 1, SYSTUR 1. Pendências: nenhuma.'),
    '15161': ('Funções: CP BACKOFFICE, Carros e Hoteis, Coordenação Suporte + Carro e Hotéis, Internacional, Não Operacional, Suporte N1, Suporte N2. Bloco “Acessos esperados”: SICA_ESFERA 2, SICA_RA 4, SYSTUR 6. Pendências: SIG 1.',
                 'Função: Carros e Hoteis. Bloco “Acessos esperados”: SICA_ESFERA 1, SICA_RA 1, SYSTUR 1. Pendências: nenhuma.'),
    '34532590': ('Funções: CP BACKOFFICE, Carros e Hoteis, Coordenação Suporte + Carro e Hotéis, Internacional, Não Operacional, Suporte N1, Suporte N2. Bloco “Acessos esperados”: SICA_ESFERA 2, SICA_RA 4, SYSTUR 6. Pendências: SIG 1.',
                 'Função: Carros e Hoteis. Bloco “Acessos esperados”: SICA_ESFERA 1, SICA_RA 1, SYSTUR 1. Pendências: nenhuma.'),
    '90000066': ('Funções: CP BACKOFFICE, Carros e Hoteis, Coordenação Suporte + Carro e Hotéis, Internacional, Não Operacional, Suporte N1, Suporte N2. Bloco “Acessos esperados”: SICA_ESFERA 2, SICA_RA 4, SYSTUR 6. Pendências: SIG 1.',
                 'Função: Carros e Hoteis. Bloco “Acessos esperados”: SICA_ESFERA 1, SICA_RA 1, SYSTUR 1. Pendências: nenhuma.'),
    '33057': ('Funções: CP BACKOFFICE, Carros e Hoteis, Coordenação Suporte + Carro e Hotéis, Internacional, Não Operacional, Suporte N1, Suporte N2. Bloco “Acessos esperados”: SICA_ESFERA 2, SICA_RA 4, SYSTUR 6. Pendências: SIG 1.',
                 'Função: Carros e Hoteis. Bloco “Acessos esperados”: SICA_ESFERA 1, SICA_RA 1, SYSTUR 1. Pendências: nenhuma.'),
    '34532095': ('Funções: CP BACKOFFICE, Carros e Hoteis, Coordenação Suporte + Carro e Hotéis, Internacional, Não Operacional, PÓS FATURAMENTO, Suporte N1, Suporte N2. Bloco “Acessos esperados”: SICA_ESFERA 2, SICA_RA 4, SYSTUR 6. Pendências: SIG 8.',
                 'Função: Internacional. Bloco “Acessos esperados”: SYSTUR 1. Pendências: SIG 1.'),
    '34532254': ('Funções: Adiantamento + Conciliação bancária, Adiantamentos a fornecedores N2, Baixa Adiantamentos a fornecedores N1, Contas a Pagar I, Pagamentos Internacionais, Pagamentos Nacionais, Pagamentos Nacionais I, Suporte ao Caixa. Bloco “Acessos esperados”: SYSTUR 8. Pendências: SIG 1.',
                 'Função: Pagamentos Nacionais I. Bloco “Acessos esperados”: SYSTUR 1. Pendências: nenhuma.'),
    '90000177': ('Funções: Atendimento a fornecedores, Atendimento a fornecedores  CVC e VISUAL, Atendimento a fornecedores  TREND, Atendimento a fornecedores  TREND I, Atendimento a fornecedores TREND N2, Atendimento ao Fornecedor - Carros, Atendimento ao Fornecedor C&H. Bloco “Acessos esperados”: SYSTUR 6. Pendências: nenhuma.',
                 'Função: Atendimento a fornecedores TREND N2. Bloco “Acessos esperados”: SYSTUR 1. Pendências: nenhuma.'),
    '90001090': ('Funções: Atendimento Backoffice - N2, Sup Financeiro, Sup Financeiro - Caixa. Bloco “Acessos esperados”: SYSTUR 3. Pendências: nenhuma.',
                 'Função: Sup Financeiro. Bloco “Acessos esperados”: nenhum. Pendências: nenhuma.'),
    '34531885': ('Funções: Atendimento Backoffice - N2, Sup Financeiro, Sup Financeiro - Caixa. Bloco “Acessos esperados”: SYSTUR 3. Pendências: nenhuma.',
                 'Função: Sup Financeiro. Bloco “Acessos esperados”: nenhum. Pendências: nenhuma.'),
}


def regra(*a, **kw):
    kw.setdefault("rot_medido", MEDIDO)
    return base.regra(*a, **kw)


def lista(doc, itens):
    for txt in itens:
        p = doc.add_paragraph(txt, style="List Bullet")
        for r_ in p.runs:
            r_.font.size = Pt(9.5)
            r_.font.color.rgb = TEXTO
        _MD.append(f"- {txt}")
    _MD.append("")


def tabela_casos(doc):
    """Um bloco por pessoa: Antes | Depois, e a linha para ela marcar."""
    for m, nome, ident, funcao, _a, _d in CASOS:
        antes, depois = ANTES_DEPOIS[m]
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(f"{m} — {nome}")
        r.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = AZUL
        r = p.add_run(f"   ·   {ident} → {funcao}")
        r.font.size = Pt(9)
        r.font.color.rgb = CINZA
        _MD.append(f"\n**{m} — {nome}** · {ident} → {funcao}\n")
        tbl = doc.add_table(rows=0, cols=2)
        tbl.style = "Table Grid"
        for rot, val in (("Antes (pacote de 28/09)", antes),
                         ("Depois (este pacote)", depois),
                         ("Confere? Se não, o que deveria ser?", "")):
            cells = tbl.add_row().cells
            cells[0].width = Cm(4.2)
            cells[1].width = Cm(12.3)
            r0 = cells[0].paragraphs[0].add_run(rot)
            r0.bold = True
            r0.font.size = Pt(8.5)
            r0.font.color.rgb = AZUL if not rot.startswith("Confere") else CINZA
            r1 = cells[1].paragraphs[0].add_run(val)
            r1.font.size = Pt(8.5)
            r1.font.color.rgb = TEXTO
            _MD.append(f"- **{rot}:** {val}")
        _MD.append("")


def main():
    _MD.clear()
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)

    # ------------------------------------------------------------- capa
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = t.add_run("CVC IAM Analytics")
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = AZUL
    s = doc.add_paragraph()
    r = s.add_run("Roteiro de validação — seu documento de 29/09 (CCO)")
    r.font.size = Pt(13)
    r.font.color.rgb = CINZA
    par(doc, "29/09/2026", size=9, cor=CINZA, md=False)
    _MD.insert(0, "# CVC IAM Analytics\n"
                  "## Roteiro de validação — seu documento de 29/09 (CCO)\n"
                  "29/09/2026\n")

    doc.add_paragraph()
    par(doc, "Por que este documento existe", bold=True, size=11, cor=AZUL, space=2)
    par(doc,
        "Ele responde o seu documento de 29/09 (ajustes_apl_29_09): na CCO, a "
        "função da pessoa passa a ser definida pelo perfil que ela tem no "
        "SYSTUR, SICA RA, SICA ESFERA ou SIGOT. Traz também os 19 casos da sua "
        "base que mudaram com isso, com o antes e o depois, para você conferir.")
    nota(doc,
         "Este pacote substitui o de 28/09 e já inclui todos os ajustes dele. Ele "
         "traz a pasta DADOS com o banco processado: você NÃO precisa rodar o "
         "Processador.")

    # ------------------------------------------------- antes de tudo
    doc.add_page_break()
    h1(doc, "Antes de tudo: o que fazer, nesta ordem")
    lista(doc, [
        "1. Feche o painel e o Processador, se estiverem abertos.",
        "2. BACKUP (obrigatório): copie as pastas DADOS\\BANCO e INTERACOES "
        "para outro lugar. Se algo não sair como esperado, é só devolvê-las.",
        "3. Extraia o pacote na pasta principal da instalação — a que contém "
        "DADOS, ENTRADA e EXECUTAVEIS — e aceite substituir os arquivos.",
        "4. NÃO apague nem mexa na pasta INTERACOES: é onde ficam as suas "
        "tratativas, e o painel as lê ao vivo.",
        "5. Abra o visualizador.exe.",
    ])

    # ------------------------------------------------- a regra
    doc.add_page_break()
    h1(doc, "1. O ajuste")

    regra(doc, "1", "A função da CCO vem do SYSTUR, SICA RA, SICA ESFERA e SIGOT",
          prioritaria=True,
          decide="Qual função da equipe (centro de custo + gestor) é cobrada de "
                 "cada pessoa na CCO.",
          criterio="Até aqui só o SYSTUR dizia a função. Quem não tinha SYSTUR "
                   "recebia TODAS as funções da equipe — daí o “faltam 7” no "
                   "Oracle e o “tem 2 de 13”. Agora vale o perfil que a pessoa "
                   "tem em qualquer um dos quatro sistemas, buscado nas linhas da "
                   "própria equipe. Oracle e SIG não definem a função (o mesmo "
                   "perfil do Oracle está em várias funções); eles só recebem o "
                   "que a função prevê. As demais funções da equipe aparecem em "
                   "“Outras funções que a pessoa pode ter”, sem cobrança. Quem "
                   "não tem perfil em nenhum dos quatro continua recebendo todas.",
          conferir="Consulta → BRUNA OLIVEIRA FERREIRA DA SILVA (34532401) → aba "
                   "Acessos. ANTES (pacote de 28/09): Oracle “Tem 1 · Faltam 7”; "
                   "“Acessos esperados (12)” com SIGOT, SIG, SICA ESFERA e SYSTUR; "
                   "“Outros acessos previstos: SICA_RA — 4 opções”; Funções "
                   "previstas (5) com “Pós Faturamento — tem 2 de 13”.",
          esperado="DEPOIS — Oracle: “Tem 1: CVC AR BRASIL Faturamento · Faltam 2: CVC AP "
                   "BRASIL Consulta, CVC AP BRASIL Cadastro de Fornecedor”. SICA "
                   "RA: POS FATURAMENTO CONC. Sem “Acessos esperados” de SIGOT, "
                   "SIG, SICA ESFERA e SYSTUR e sem “Outros acessos previstos”. "
                   "Funções previstas: Pós Faturamento — “tem 2 de 4”. Outras "
                   "funções que a pessoa pode ter: as outras 5 da equipe.")

    # ------------------------------------------------- casos
    doc.add_page_break()
    h1(doc, "2. Os 19 casos da sua base que mudaram")
    par(doc,
        "Todos são da CCO, não têm SYSTUR e têm perfil em SICA RA ou SIGOT. Antes "
        "recebiam todas as funções da equipe; agora recebem só a função do perfil "
        "que têm. Para cada um: Consulta → busque a matrícula → aba Acessos, e "
        "compare com o “Depois”. As funções que saíram continuam visíveis em "
        "“Outras funções que a pessoa pode ter”, sem cobrança.")
    tabela_casos(doc)
    par(doc, "Destaques", bold=True, size=10.5, cor=AZUL, space=2)
    lista(doc, [
        "ERICKA CRISTINA REIS (12698): no ajuste de 28/09 ela continuou em análise "
        "porque tinha 3 dos 8 perfis de SIG previstos. Os 8 eram a soma de duas "
        "funções. Com a função certa (Carros e Hoteis), os 3 que ela tem são "
        "exatamente os previstos e o SIG fica aderente.",
        "KATHELIN LOPES CORREA (34532095): das 8 pendências sobra 1 — o SIG "
        "COM_INFORMATIVOS, que a função Internacional não prevê (“acesso fora da "
        "função”).",
        "BRUNA OLIVEIRA, MICHELE, ANA CAROLINE, PALOMA CAMPOREZE, ALYNE, JESSICA "
        "LOPES e PALOMA ISABEL são da mesma equipe (01.06.02.01): cada uma fica "
        "com a sua função de Pós Faturamento / Pós Terrestre.",
    ])

    # ------------------------------------------------- o que muda na tela
    doc.add_page_break()
    h1(doc, "3. O que vai parecer diferente")
    lista(doc, [
        "O total de linhas cai de 13.446 para 13.114: são os esperados de outras "
        "funções que essas 19 pessoas recebiam.",
        "As pendências caem de 1.072 para 1.064 pessoas (de 1.123 para 1.108 "
        "linhas). Nas 19 pessoas, de 16 linhas para 1.",
        "Nenhuma outra pessoa muda: quem tem SYSTUR já tinha a função definida, e "
        "quem não é da CCO não passa por esta regra.",
    ])

    # ------------------------------------------------- pergunta
    h1(doc, "4. Uma pergunta para você")
    par(doc,
        "No bloco “Sem mapeamento”, os sistemas que a função da pessoa não prevê "
        "aparecem com o texto “sem perfil previsto para o cargo/centro de custo”. "
        "Para quem é da CCO, prefere que diga “a função da pessoa não prevê "
        "acesso a este sistema”?", size=9.5)
    par(doc, "Sua resposta: ______________________________________________",
        size=9.5, cor=CINZA)

    # ------------------------------------------------- saida
    OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_DOCX)
    OUT_MD.write_text("\n".join(_MD), encoding="utf-8")
    print(f"gerado: {OUT_DOCX}")
    print(f"gerado: {OUT_MD}")


if __name__ == "__main__":
    main()
