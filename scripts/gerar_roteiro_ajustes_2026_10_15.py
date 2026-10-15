# -*- coding: utf-8 -*-
"""Gera ENTREGA/ROTEIRO_AJUSTES_CVC_IAM_2026-10-15.docx (+ .md).

Atualizacao de 15/10 (pacote versao 1.0.3): aba Transferidos com o resumo por
sistema (mockup aprovado pela Bruna em 15/10). Mesma formatacao dos roteiros
anteriores.

Uso:  python scripts/gerar_roteiro_ajustes_2026_10_15.py
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import gerar_roteiro_ajustes as base  # noqa: E402

from docx import Document  # noqa: E402
from docx.shared import Pt  # noqa: E402

h1, par = base.h1, base.par
AZUL, CINZA, TEXTO = base.AZUL, base.CINZA, base.TEXTO
_MD = base._MD

RAIZ = Path(__file__).resolve().parent.parent
OUT_DOCX = RAIZ / "ENTREGA" / "ROTEIRO_AJUSTES_CVC_IAM_2026-10-15.docx"
OUT_MD = RAIZ / "ENTREGA" / "ROTEIRO_AJUSTES_CVC_IAM_2026-10-15.md"


def lista(doc, itens, estilo="List Number"):
    for txt in itens:
        p = doc.add_paragraph(txt, style=estilo)
        for r_ in p.runs:
            r_.font.size = Pt(9.5)
            r_.font.color.rgb = TEXTO
        _MD.append(f"- {txt}")
    _MD.append("")


def main():
    _MD.clear()
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)

    t = doc.add_paragraph()
    r = t.add_run("CVC IAM Analytics")
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = AZUL
    s = doc.add_paragraph()
    r = s.add_run("Atualização de 15/10 — resumo por sistema na aba Transferidos")
    r.font.size = Pt(13)
    r.font.color.rgb = CINZA
    par(doc, "15/10/2026", size=9, cor=CINZA, md=False)
    _MD.insert(0, "# CVC IAM Analytics\n"
                  "## Atualização de 15/10 — resumo por sistema na aba Transferidos\n"
                  "15/10/2026\n")

    h1(doc, "1. O que muda")
    par(doc,
        "Na aba Transferidos, ao abrir uma pessoa, a tabela de 4 colunas (Sistema, "
        "Tem atualmente, Mapeando nova área, Incluir/excluir/alterar) foi trocada "
        "pelo resumo por sistema do mockup aprovado em 15/10, no mesmo padrão da "
        "tela de Consulta.")
    lista(doc, [
        "Aderente: a pessoa já tem um perfil permitido na nova área. Nada a fazer. "
        "Se a nova área prevê outras opções no mesmo sistema, elas aparecem como "
        "“pode ter também”, sem precisar incluir.",
        "Trocar: a pessoa não tem o perfil permitido. A linha diz qual perfil "
        "excluir e qual incluir. Quando há mais de uma opção, aparece “incluir um "
        "de: A ou B”.",
        "Incluir: a pessoa não tem acesso ao sistema e a nova área prevê.",
        "Excluir: a pessoa tem acesso, mas a nova área não prevê esse sistema.",
        "Sem mapeamento: a nova área não prevê nada no sistema. Nada a fazer.",
        "Sem extrato: a nova área prevê o Opera Operacional, que não tem extrato. "
        "Aparece só para informação e não conta como ação.",
    ], estilo="List Bullet")
    lista(doc, [
        "No topo do resumo aparece o veredito: “Aderente na nova área, nada a "
        "alterar” ou “N ações a fazer”.",
        "Em “ver detalhe” ficam 4 abas: Tem hoje (com login), Previsto na nova "
        "área, Sem mapeamento e Funções previstas.",
        "Na lista, ao lado do número de acessos, o selo mostra “N ações” ou "
        "“aderente”.",
        "No Oracle e no SIG cada perfil é um acesso. Por isso, o que falta "
        "continua como inclusão, mesmo que a pessoa já tenha outro perfil.",
        "Listas com mais de 4 perfis aparecem recolhidas (“41 perfis +”). Clique "
        "para abrir.",
        "No Excel da aba, a outra opção de um sistema em que a pessoa já é "
        "aderente sai como “pode ter também (não precisa incluir)”, e o acesso "
        "fora do padrão das duas áreas sai como “fora do padrão”.",
        "A situação Revisar/Tratado e os números do topo da aba não mudaram.",
    ], estilo="List Bullet")
    par(doc,
        "Este pacote também traz as correções anteriores (banco corrompido, trava "
        "de execução pela rede e o aviso “há uma carga mais recente” que não "
        "sumia).", size=9.5)

    h1(doc, "2. Instalação na rede")
    lista(doc, [
        "Feche o painel e o Processador em todas as máquinas.",
        "Copie a pasta EXECUTAVEIS do zip por cima da EXECUTAVEIS da rede.",
        "Não é preciso reprocessar: o painel novo lê o banco que já está na rede.",
        "Abra o painel pela sua cópia local (ex.: C:\\CVC_IAM\\EXECUTAVEIS\\"
        "visualizador.exe). Como a versão mudou para 1.0.3, ele se atualiza "
        "sozinho a partir da rede.",
    ])
    par(doc,
        "Se o painel não se atualizar sozinho, copie a EXECUTAVEIS da rede por "
        "cima da sua cópia local e abra de novo.", size=9.5)

    h1(doc, "3. O que conferir")
    lista(doc, [
        "Na aba Transferidos, abra uma pessoa com o “+”: aparece o “Resumo por "
        "sistema” com uma situação por sistema e o veredito no topo.",
        "Uma pessoa que já tem o perfil permitido num sistema (ex.: IC_CADASTRO) "
        "aparece como Aderente, com as outras opções em “pode ter também”, sem "
        "“Incluir”.",
        "“ver detalhe” abre as 4 abas, e a aba Tem hoje mostra o login de cada "
        "sistema.",
        "O Excel da aba sai com “pode ter também (não precisa incluir)” nesses "
        "casos.",
        "Os números do topo da aba (a revisar, tratados, total) são os mesmos de "
        "antes da atualização.",
    ], estilo="List Bullet")

    OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_DOCX)
    OUT_MD.write_text("\n".join(_MD), encoding="utf-8")
    print(f"gerado: {OUT_DOCX}")
    print(f"gerado: {OUT_MD}")


if __name__ == "__main__":
    main()
