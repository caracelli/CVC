# -*- coding: utf-8 -*-
"""Gera ENTREGA/ROTEIRO_AJUSTES_CVC_IAM_2026-09-18.docx (+ .md).

Responde o retorno da area de 17/09 ("Aplicação_17_09.docx", 4 pedidos) e o
pedido que veio por mensagem em 18/09 (contador da aba Bases). Cobre tambem,
resumido, a rodada de 16/09 (os 23 nomes que "nao vinham na aplicacao"), porque
o pacote entregue leva as duas.

Mesma formatacao dos roteiros de 08/09 e 11/09 (scripts/gerar_roteiro_ajustes.py).

OS NUMEROS SAO DA BASE DELA: o DADOS.zip de 15/09 reprocessado aqui com o codigo
desta entrega. Cada regra diz se o valor e' INVARIANTE (vale em qualquer base) ou
MEDIDO (pode diferir se a ENTRADA dela mudou).

Uso:  python scripts/gerar_roteiro_ajustes_2026_09_18.py
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import gerar_roteiro_ajustes as base  # noqa: E402

from docx import Document  # noqa: E402
from docx.shared import Pt  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402

h1, par, nota = base.h1, base.par, base.nota
AZUL, CINZA, TEXTO = base.AZUL, base.CINZA, base.TEXTO
_MD = base._MD
MEDIDO = "Medido na sua base (a de 15/09, reprocessada aqui com esta versão)"

RAIZ = Path(__file__).resolve().parent.parent
OUT_DOCX = RAIZ / "ENTREGA" / "ROTEIRO_AJUSTES_CVC_IAM_2026-09-18.docx"
OUT_MD = RAIZ / "ENTREGA" / "ROTEIRO_AJUSTES_CVC_IAM_2026-09-18.md"


def regra(*a, **kw):
    kw.setdefault("rot_medido", MEDIDO)
    return base.regra(*a, **kw)


def pergunta(doc, titulo, texto):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(titulo)
    r.bold = True
    r.font.size = Pt(10.5)
    r.font.color.rgb = AZUL
    _MD.append(f"\n### {titulo}\n")
    par(doc, texto, size=9.5)
    par(doc, "Sua resposta: ______________________________________________",
        size=9.5, cor=CINZA)


def lista(doc, itens):
    for txt in itens:
        p = doc.add_paragraph(txt, style="List Bullet")
        for r_ in p.runs:
            r_.font.size = Pt(9.5)
            r_.font.color.rgb = TEXTO
        _MD.append(f"- {txt}")
    _MD.append("")


def main():
    _MD.clear()
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)

    # ------------------------------------------------------------- capa
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = t.add_run("CVC IAM Analytics")
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = AZUL
    s = doc.add_paragraph()
    r = s.add_run("Roteiro de validação — o que você pediu em 17 e 18/09")
    r.font.size = Pt(13)
    r.font.color.rgb = CINZA
    par(doc, "18/09/2026", size=9, cor=CINZA, md=False)
    _MD.insert(0, "# CVC IAM Analytics\n"
                  "## Roteiro de validação — o que você pediu em 17 e 18/09\n"
                  "18/09/2026\n")

    doc.add_paragraph()
    par(doc, "Por que este documento existe", bold=True, size=11, cor=AZUL, space=2)
    par(doc,
        "Ele responde, um a um, os quatro pontos do seu documento de 17/09 e o "
        "pedido do contador de registros que você mandou em 18/09. No fim há "
        "quatro perguntas — são decisões suas, e a aplicação segue como está até "
        "você responder.")
    par(doc, "Como usar", bold=True, size=11, cor=AZUL, space=2)
    par(doc,
        "Igual aos roteiros de 08/09 e 11/09: cada ajuste tem “Como conferir” e o "
        "valor esperado, e a última linha pergunta se a regra está CERTA. As "
        "marcadas com ★ são as que mais aparecem na tela.")
    nota(doc,
         "Se o pacote que você recebeu tiver a pasta DADOS, o banco já vai "
         "processado e você NÃO precisa rodar o Processador. Se ele tiver só "
         "EXECUTAVEIS, rode o Processador UMA vez depois de copiar: o item 1 "
         "(funções) e os acessos a incluir dependem dessa rodada.")

    # ------------------------------------------------- antes de tudo
    doc.add_page_break()
    h1(doc, "Antes de tudo: o que fazer, nesta ordem")
    lista(doc, [
        "1. Feche o painel e o Processador, se estiverem abertos.",
        "2. BACKUP (obrigatório): copie as pastas DADOS\\BANCO e INTERACOES para "
        "outro lugar. Se algo não sair como esperado, é só devolvê-las.",
        "3. Extraia o pacote na pasta principal da instalação — a que contém "
        "DADOS, ENTRADA e EXECUTAVEIS — e aceite substituir os arquivos.",
        "4. NÃO apague nem mexa na pasta INTERACOES: é onde ficam as suas "
        "tratativas, e o painel as lê ao vivo.",
        "5. Se o pacote NÃO trouxer a pasta DADOS, rode o Processador.exe uma vez "
        "e espere terminar (alguns minutos; não feche a janela no meio).",
        "6. Abra o visualizador.exe.",
    ])

    # ------------------------------------------------- os 5 pedidos
    doc.add_page_break()
    h1(doc, "1. O que você pediu em 17/09 e 18/09")

    regra(doc, "1.1", "Funções do Mapeamento CCO, com abrir e fechar",
          prioritaria=True,
          decide="De onde vem o acesso esperado da pessoa: matriz por cargo ou "
                 "Mapeamento CCO (centro de custo + gestor).",
          criterio="Quando o esperado vem do CCO, a aplicação passa a guardar "
                   "também a FUNÇÃO daquele acesso. Na Consulta, abaixo dos blocos "
                   "de acessos, aparece a lista “Funções previstas”, recolhida. "
                   "Cada função abre no “+” e mostra os acessos que a formam, com "
                   "sistema, perfil e se a pessoa tem ou falta. As funções com algo "
                   "faltando vêm primeiro; a que está completa aparece marcada.",
          conferir="Consulta → busque uma pessoa do centro de custo 01.06.04.01 "
                   "(por exemplo BRENDA SILVA BRITO) → “ver detalhe” → role até "
                   "“Funções previstas”.",
          esperado="7 funções. Ao abrir “Atendimento a fornecedores  CVC e VISUAL”: "
                   "“tem 1 de 5” — SYSTUR ATD_FOR_CVC_VISUAL_CP (tem) e os quatro "
                   "CVC AP … Consulta do Oracle EBS (faltam).",
          tipo_esperado="MEDIDO")

    regra(doc, "1.2", "Quem não tem mapeamento deixa de aparecer como “Aderente”",
          prioritaria=True,
          decide="O status que a Consulta mostra para quem a matriz não prevê nada.",
          criterio="Aderência é por SISTEMA. Quem não tem nenhum acesso aderente "
                   "passa a aparecer como “Sem perfis mapeados”, em cinza, e não "
                   "mais como “Aderente”. Quem é aderente em um sistema continua "
                   "“Aderente”, mesmo sem mapeamento nos outros. Continua sendo "
                   "informativo: não entra em Pendências nem nos cards.",
          conferir="Consulta → coluna Status. Procure alguém com “sem mapeamento” "
                   "no detalhe (o seu exemplo era PAULO HENRIQUE FICUCHIELLO).",
          esperado="Status “Sem perfis mapeados”: 589 linhas assim, sendo 539 "
                   "pessoas que só têm essa linha.",
          tipo_esperado="MEDIDO")

    regra(doc, "1.3", "Aviso de mais de um perfil no SYSTUR",
          prioritaria=True,
          decide="Quantos perfis a pessoa tem no mesmo sistema.",
          criterio="No SYSTUR, quando a pessoa tem mais de um perfil, a Consulta "
                   "mostra o aviso “Usuário com mais de um acesso”. É um aviso: não "
                   "vira pendência, seguindo o que já valia para perfil a mais "
                   "(a cobrança depende de uma configuração, hoje desligada). Vale "
                   "só para o SYSTUR — no SIG, ter vários perfis é o normal.",
          conferir="Consulta → busque THAIANE TAVARES SILVA DE PAULA (matrícula "
                   "14510) ou ADRIANA TATEISHI (1243) → “ver detalhe”.",
          esperado="O aviso aparece acima dos perfis. São 68 pessoas com mais de "
                   "um perfil no SYSTUR: 62 com dois, 1 com três e 5 com 42.",
          tipo_esperado="MEDIDO")

    regra(doc, "1.4", "Transferidos: o que tem, o que deveria ter, o que alterar",
          prioritaria=True,
          decide="A leitura da aba Transferidos e da planilha exportada.",
          criterio="Ao abrir a pessoa, o detalhe passa a ter três blocos: 1) os "
                   "acessos que ela tem; 2) os que deveria ter na nova área; 3) o "
                   "que precisa alterar, separando REVOGAR (sobrou da área anterior) "
                   "de INCLUIR (a nova equipe tem e ela não), com a função de cada "
                   "acesso. Quando não há diferença, aparece “Validado: os acessos "
                   "da nova área já estão aderentes”. A planilha ganhou as colunas "
                   "“Situação do acesso” e “Função (CCO)” e agora também traz as "
                   "linhas do que falta incluir. As colunas antigas não mudaram de "
                   "lugar.",
          conferir="Transferidos → abra uma pessoa no “+” → veja os três blocos. "
                   "Depois clique em Exportar Excel e confira as duas colunas novas, "
                   "no fim da planilha.",
          esperado="199 pessoas na aba e 295 acessos já com a função preenchida.",
          tipo_esperado="MEDIDO")

    regra(doc, "1.5", "Contador de registros: arquivo × aplicação",
          prioritaria=True,
          decide="Se tudo o que estava no arquivo entrou na aplicação.",
          criterio="A tela de bases (o link “Arquivos importados”) passa a mostrar, "
                   "em cada base, quantas linhas o arquivo tinha e quantas estão na "
                   "aplicação, com “confere” quando bate. Para RH e diretório AD a "
                   "tela diz “acumulado”: essas bases somam as cargas anteriores por "
                   "desenho, então o total é maior que o do arquivo e comparar um a "
                   "um acusaria um erro que não existe.",
          conferir="Clique em “Arquivos importados”, abra um setor e olhe a linha "
                   "abaixo do nome da base.",
          esperado="Todos os extratos e as matrizes conferem: SYSTUR "
                   "7.074, SIG 80.140, SIGOT 208, Oracle EBS 2.804, SICA RA 257, "
                   "SICA Esfera 44, IC 65, matriz de perfis 2.699 e CCO 1.241.",
          tipo_esperado="MEDIDO")

    # ------------------------------------------------- rodada anterior
    doc.add_page_break()
    h1(doc, "2. O que já tinha entrado na rodada de 16/09")
    par(doc,
        "Se você ainda não aplicou o pacote de 16/09, ele está incluído aqui. "
        "Resumo do que muda:")
    lista(doc, [
        "Os 23 nomes que você listou como “ativos que seguem não vindo na "
        "aplicação” voltaram a aparecer na Consulta — os 23.",
        "Quem não tem expectativa de acesso aparece como “Não Mapeado”, com o "
        "texto “Não tem mapeamento localizado para o centro de custo.”.",
        "Quem saiu do arquivo de ativos mais recente deixou de aparecer como "
        "ativo (eram pessoas que já constavam na base de desligados).",
        "A regra de provável desligamento passou a valer só para quem também "
        "sumiu do arquivo de ativos. Com isso, 318 pessoas que estavam ativas e "
        "tinham perdido o acesso voltaram a aparecer como “Incluir Acesso”.",
    ])

    # ------------------------------------------------- perguntas
    doc.add_page_break()
    h1(doc, "3. Quatro perguntas para você")
    pergunta(doc, "3.1  A matriz do centro de custo 05.12.02.01 mudou de propósito?",
             "Na matriz do SYSTUR de 18/08 saíram as linhas de GERENTE ATENDIMENTO "
             "e GERENTE DE OPERAÇÕES desse centro de custo; ficou só GERENTE "
             "EXECUTIVO ATENDIMENTO. É por isso que a equipe da PATRICIA MARAGNA "
             "aparece sem mapeamento. Foi intencional, ou essas linhas deveriam "
             "voltar?")
    pergunta(doc, "3.2  O limiar de 30% deve continuar cortando esses casos?",
             "Quando menos de 30% das pessoas do cargo têm um acesso, a aplicação "
             "não cobra a inclusão. Na sua lista de 23 nomes, 9 estão nessa "
             "situação: têm perfil previsto na matriz, mas o acesso não é cobrado. "
             "Quer que passem a aparecer como “Incluir Acesso”?")
    pergunta(doc, "3.3  As contas com 42 perfis no SYSTUR entram no aviso?",
             "Das 68 pessoas com mais de um perfil no SYSTUR, 5 têm 42 perfis e "
             "parecem contas técnicas. Elas devem receber o mesmo aviso, ou ficam "
             "de fora? E o aviso deve continuar informativo, ou virar pendência?")
    pergunta(doc, "3.4  Por que a mesma matrícula aparece duas vezes nos desligados?",
             "O arquivo de desligados de 15/09 tem 11.634 linhas, mas 11.072 "
             "matrículas distintas: 556 matrículas aparecem repetidas. A aplicação "
             "considera cada pessoa uma vez, então nada se perdeu — mas vale "
             "entender se é rescisão dupla, vínculo antigo ou erro de extração.")

    # ------------------------------------------------- correções
    doc.add_page_break()
    h1(doc, "4. Duas correções em documentos anteriores")
    lista(doc, [
        "Perfil a mais (perfil excessivo): o roteiro de 08/09 fala em 141 casos e "
        "178 perfis; o número correto, medido nos 7 sistemas, é 130 casos e 186 "
        "perfis a mais.",
        "O roteiro de 06/08 diz, na seção “O que ainda não tem regra”, que o perfil "
        "a mais não é sinalizado. Isso deixou de valer em 28/08: ele aparece na "
        "Consulta (“N a mais”), e só a cobrança continua desligada.",
    ])

    h1(doc, "5. O que NÃO mudou")
    par(doc,
        "Nenhuma regra de validação foi alterada nesta entrega. Comparando a sua "
        "base antes e depois, os números de Aderente (6.344), Em Análise (508) e "
        "Alterar Perfil (318) ficaram idênticos, assim como as divergências de "
        "desligados, transferidos e contas de serviço, o RH, os acessos e o ciclo "
        "de vida. Seguem valendo as decisões já fechadas: franqueado sem acesso não "
        "recebe sugestão, prestador sem grupo de espelho não gera linha, e a "
        "cobrança de perfil a mais continua desligada.")

    doc.save(OUT_DOCX)
    OUT_MD.write_text("\n".join(_MD), encoding="utf-8")
    print(f"OK -> {OUT_DOCX.name} ({OUT_DOCX.stat().st_size/1024:.0f} KB)")
    print(f"OK -> {OUT_MD.name} ({OUT_MD.stat().st_size/1024:.0f} KB)")


if __name__ == "__main__":
    main()
