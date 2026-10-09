# -*- coding: utf-8 -*-
"""Gera ENTREGA/ROTEIRO_AJUSTES_CVC_IAM_2026-10-09.docx (+ .md).

Atualizacao de 09/10 (pacote versao 1.0.2): "database disk image is malformed" no painel da rede.
Mesma formatacao dos roteiros anteriores.

Uso:  python scripts/gerar_roteiro_ajustes_2026_10_09.py
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import gerar_roteiro_ajustes as base  # noqa: E402

from docx import Document  # noqa: E402
from docx.shared import Pt  # noqa: E402

h1, par = base.h1, base.par
AZUL, CINZA, TEXTO = base.AZUL, base.CINZA, base.TEXTO
_MD = base._MD

RAIZ = Path(__file__).resolve().parent.parent
OUT_DOCX = RAIZ / "ENTREGA" / "ROTEIRO_AJUSTES_CVC_IAM_2026-10-09.docx"
OUT_MD = RAIZ / "ENTREGA" / "ROTEIRO_AJUSTES_CVC_IAM_2026-10-09.md"


def lista(doc, itens, estilo="List Number"):
    for txt in itens:
        p = doc.add_paragraph(txt, style=estilo)
        for r_ in p.runs:
            r_.font.size = Pt(9.5)
            r_.font.color.rgb = TEXTO
        _MD.append(f"- {txt}")
    _MD.append("")


def main():
    _MD.clear()
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)

    t = doc.add_paragraph()
    r = t.add_run("CVC IAM Analytics")
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = AZUL
    s = doc.add_paragraph()
    r = s.add_run("Atualização de 09/10 — banco corrompido e execução pela rede")
    r.font.size = Pt(13)
    r.font.color.rgb = CINZA
    par(doc, "09/10/2026", size=9, cor=CINZA, md=False)
    _MD.insert(0, "# CVC IAM Analytics\n"
                  "## Atualização de 09/10 — banco corrompido e execução pela rede\n"
                  "09/10/2026\n")

    h1(doc, "1. O que aconteceu")
    par(doc,
        "O banco da rede (iam_analytics.db) estava íntegro, com o processamento "
        "das 10:59. Ao lado dele ficou um arquivo iam_analytics.db-wal, resto de "
        "uma gravação anterior. Ao copiar o banco para a máquina, o painel "
        "aplicava esse resto por cima do banco novo e a cópia saía corrompida. "
        "Nenhum dado nem tratativa se perdeu.")
    par(doc,
        "A causa é o modo de gravação “WAL” do SQLite, que não é seguro em pasta "
        "de rede. Esta atualização deixa de usá-lo.")

    h1(doc, "2. O que muda nesta atualização")
    lista(doc, [
        "Processador: grava o banco da rede no modo seguro para pasta de rede. "
        "Se encontrar um -wal órfão, guarda-o como .orfao_<data> (não apaga).",
        "Painel: copia só o arquivo do banco (ignora -wal ao lado) e confere a "
        "cópia antes de usar. Se a cópia estiver ruim, mantém a anterior e avisa.",
        "Painel: descarta sozinho uma cópia local corrompida.",
        "Painel: não copia enquanto o Processador está rodando. Mostra "
        "“Processamento em andamento na rede” e só oferece a carga nova quando "
        "ele termina.",
        "Painel e Processador não abrem mais direto da pasta de rede: mostram "
        "um aviso. Rodado da rede, o programa prende a pasta (ninguém consegue "
        "atualizar) e o painel passa a usar um único arquivo de cópia na rede "
        "para todos, o que corrompe o banco. Cada pessoa usa a sua cópia local "
        "de EXECUTAVEIS (ex.: C:\\CVC_IAM\\EXECUTAVEIS); os dados continuam na rede.",
    ], estilo="List Bullet")

    h1(doc, "3. Instalação na rede")
    lista(doc, [
        "Feche o painel e o Processador em todas as máquinas. Quem abriu da "
        "rede: no Gerenciador de Tarefas, finalize launcher_visualizador.exe e "
        "launcher_processador.exe (o painel segue rodando após fechar o navegador). "
        "Sem saber quem abriu: renomeie os .exe da EXECUTAVEIS da rede (e da "
        "subpasta launcher) para .exe.old_0910 — um exe em uso pode ser "
        "renomeado, só não sobrescrito. Apague os .old_0910 depois.",
        "Em DADOS\\BANCO: se ainda existirem iam_analytics.db-wal e "
        "iam_analytics.db-shm, mova os dois para uma pasta _orfao_0910 (não "
        "apague; não mexa no iam_analytics.db).",
        "Copie a pasta EXECUTAVEIS do zip por cima da EXECUTAVEIS da rede.",
        "No seu computador, copie a EXECUTAVEIS da rede (já atualizada) por cima "
        "da sua cópia local (ex.: C:\\CVC_IAM\\EXECUTAVEIS) e abra o Processador.exe "
        "de lá. Uma máquina só, com todos os painéis fechados; espere "
        "“Processamento finalizado”.",
        "Abra o painel (visualizador.exe da cópia local).",
    ])
    par(doc,
        "Versão: o pacote vem como 1.0.2. As máquinas só se atualizam sozinhas "
        "quando a versão do config da rede muda. Enquanto ela ficar 1.0.2, quem "
        "usar o programa precisa copiar a EXECUTAVEIS nova para a máquina.",
        size=9.5)

    h1(doc, "4. O que conferir")
    lista(doc, [
        "O painel abre sem a mensagem “database disk image is malformed”.",
        "Em DADOS\\BANCO da rede ficam só o iam_analytics.db (sem -wal/-shm) "
        "depois do processamento.",
        "Com o Processador rodando, o botão Atualizar responde “Processamento em "
        "andamento na rede”; depois que ele termina, a carga nova aparece.",
        "Abrir o visualizador.exe ou o Processador.exe direto da pasta de rede "
        "mostra o aviso “não pode ser aberto direto da pasta de rede” e não abre.",
        "Os números batem com o processamento anterior (ex.: franqueados pela "
        "“Matriz franqueado”, CLAUDIA aderente pela CCO).",
    ], estilo="List Bullet")

    OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_DOCX)
    OUT_MD.write_text("\n".join(_MD), encoding="utf-8")
    print(f"gerado: {OUT_DOCX}")
    print(f"gerado: {OUT_MD}")


if __name__ == "__main__":
    main()
